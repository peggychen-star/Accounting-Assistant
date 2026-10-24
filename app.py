"""Preview, filter, hide columns, build pivot summaries, and export Excel."""

import copy
import hashlib
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st


st.set_page_config(page_title="Accounting Excel Preview", page_icon="📊", layout="wide")


def reset_workspace(source_id):
    """Clear edits when the uploaded workbook or worksheet changes."""
    for key in list(st.session_state.keys()):
        if key.startswith(("filter_column_", "filter_values_", "delete_columns_", "pivot_category_", "pivot_value_", "pivot_split_", "pivot_split_values_")):
            del st.session_state[key]

    st.session_state.filter_source_id = source_id
    st.session_state.filters = []
    st.session_state.delete_tabs = []
    st.session_state.pivot_tabs = []
    st.session_state.function_tabs = []
    st.session_state.filter_history = []
    st.session_state.next_filter_id = 0
    st.session_state.next_delete_id = 0
    st.session_state.next_pivot_id = 0


def prepare_workspace():
    """Preserve filters if the app is updated while Streamlit is running."""
    if "delete_tabs" not in st.session_state:
        st.session_state.delete_tabs = []
    if "function_tabs" not in st.session_state:
        st.session_state.function_tabs = [
            {"kind": "filter", "id": item["id"]}
            for item in st.session_state.get("filters", [])
        ]
    if "next_delete_id" not in st.session_state:
        st.session_state.next_delete_id = 0
    if "pivot_tabs" not in st.session_state:
        st.session_state.pivot_tabs = []
    if "next_pivot_id" not in st.session_state:
        st.session_state.next_pivot_id = 0

    # Existing sessions from previous app versions did not contain split settings.
    for item in st.session_state.pivot_tabs:
        item.setdefault("split", None)
        item.setdefault("split_selected", [])
        pivot_id = item["id"]
        st.session_state.setdefault(f"pivot_split_{pivot_id}", item["split"])
        st.session_state.setdefault(
            f"pivot_split_values_{pivot_id}", list(item["split_selected"])
        )


def save_undo_step():
    """Save filter, delete-column, and pivot edits for the Back button."""
    st.session_state.filter_history.append(
        copy.deepcopy(
            {
                "filters": st.session_state.filters,
                "delete_tabs": st.session_state.delete_tabs,
                "pivot_tabs": st.session_state.pivot_tabs,
                "function_tabs": st.session_state.function_tabs,
            }
        )
    )


def add_filter():
    """Add an independent filter tab."""
    save_undo_step()
    filter_id = st.session_state.next_filter_id
    st.session_state.next_filter_id += 1
    st.session_state.filters.append({"id": filter_id, "column": None, "selected": []})
    st.session_state.function_tabs.append({"kind": "filter", "id": filter_id})
    st.session_state[f"filter_column_{filter_id}"] = None
    st.session_state[f"filter_values_{filter_id}"] = []


def add_delete_column():
    """Add a tab for removing columns from the preview and export."""
    save_undo_step()
    delete_id = st.session_state.next_delete_id
    st.session_state.next_delete_id += 1
    st.session_state.delete_tabs.append({"id": delete_id, "columns": []})
    st.session_state.function_tabs.append({"kind": "delete", "id": delete_id})
    st.session_state[f"delete_columns_{delete_id}"] = []


def add_pivot():
    """Start a pivot setup without changing any existing filters/deletions."""
    save_undo_step()
    pivot_id = st.session_state.next_pivot_id
    st.session_state.next_pivot_id += 1
    st.session_state.pivot_tabs.append(
        {"id": pivot_id, "category": None, "value": None,
         "split": None, "split_selected": []}
    )
    st.session_state[f"pivot_category_{pivot_id}"] = None
    st.session_state[f"pivot_value_{pivot_id}"] = None
    st.session_state[f"pivot_split_{pivot_id}"] = None
    st.session_state[f"pivot_split_values_{pivot_id}"] = []


def change_pivot_category(pivot_id):
    """Keep a pivot's category and sum column independent from other pivots."""
    new_category = st.session_state[f"pivot_category_{pivot_id}"]
    for item in st.session_state.pivot_tabs:
        if item["id"] == pivot_id and item["category"] != new_category:
            save_undo_step()
            item["category"] = new_category
            # A column cannot simultaneously be both the group and sum column.
            if item["value"] == new_category:
                item["value"] = None
                st.session_state[f"pivot_value_{pivot_id}"] = None
            if item.get("split") == new_category:
                item["split"] = None
                item["split_selected"] = []
                st.session_state[f"pivot_split_{pivot_id}"] = None
                st.session_state[f"pivot_split_values_{pivot_id}"] = []
            break


def change_pivot_value(pivot_id):
    new_value = st.session_state[f"pivot_value_{pivot_id}"]
    for item in st.session_state.pivot_tabs:
        if item["id"] == pivot_id and item["value"] != new_value:
            save_undo_step()
            item["value"] = new_value
            if item.get("split") == new_value:
                item["split"] = None
                item["split_selected"] = []
                st.session_state[f"pivot_split_{pivot_id}"] = None
                st.session_state[f"pivot_split_values_{pivot_id}"] = []
            break


def change_pivot_split(pivot_id):
    """Record the column used to spread pivot totals across new columns."""
    chosen = st.session_state[f"pivot_split_{pivot_id}"]
    for item in st.session_state.pivot_tabs:
        if item["id"] == pivot_id and item.get("split") != chosen:
            save_undo_step()
            item["split"] = chosen
            item["split_selected"] = []
            st.session_state[f"pivot_split_values_{pivot_id}"] = []
            break


def change_pivot_split_values(pivot_id):
    """Preserve multi-selected split categories for undo and export."""
    chosen = list(st.session_state[f"pivot_split_values_{pivot_id}"])
    for item in st.session_state.pivot_tabs:
        if item["id"] == pivot_id and item.get("split_selected", []) != chosen:
            save_undo_step()
            item["split_selected"] = chosen
            break


def change_column(filter_id):
    """Record a filter's column change and reset its selected values."""
    selected_column = st.session_state[f"filter_column_{filter_id}"]
    for item in st.session_state.filters:
        if item["id"] == filter_id:
            if item["column"] != selected_column:
                save_undo_step()
                item["column"] = selected_column
                item["selected"] = []
                st.session_state[f"filter_values_{filter_id}"] = []
            break


def change_values(filter_id):
    """Remember changes to a filter's multiselect for the Back button."""
    chosen = list(st.session_state[f"filter_values_{filter_id}"])
    for item in st.session_state.filters:
        if item["id"] == filter_id:
            if item["selected"] != chosen:
                save_undo_step()
                item["selected"] = chosen
            break


def change_deleted_columns(delete_id):
    """Remember changes to columns selected for deletion from the preview."""
    chosen = list(st.session_state[f"delete_columns_{delete_id}"])
    for item in st.session_state.delete_tabs:
        if item["id"] == delete_id:
            if item["columns"] != chosen:
                save_undo_step()
                item["columns"] = chosen
            break


def go_back():
    """Undo the latest filter, delete-column, or pivot action."""
    if not st.session_state.filter_history:
        return

    previous = st.session_state.filter_history.pop()
    # Backward compatibility with history created by the earlier filter-only app.
    if isinstance(previous, list):
        previous = {
            "filters": previous,
            "delete_tabs": [],
            "function_tabs": [
                {"kind": "filter", "id": item["id"]} for item in previous
            ],
            "pivot_tabs": [],
        }

    st.session_state.filters = previous["filters"]
    st.session_state.delete_tabs = previous["delete_tabs"]
    st.session_state.pivot_tabs = previous.get("pivot_tabs", [])
    st.session_state.function_tabs = previous["function_tabs"]

    # Restore widget values before the widgets are created on the next rerun.
    for key in list(st.session_state.keys()):
        if key.startswith(("filter_column_", "filter_values_", "delete_columns_", "pivot_category_", "pivot_value_", "pivot_split_", "pivot_split_values_")):
            del st.session_state[key]

    for item in st.session_state.filters:
        st.session_state[f"filter_column_{item['id']}"] = item["column"]
        st.session_state[f"filter_values_{item['id']}"] = list(item["selected"])
    for item in st.session_state.delete_tabs:
        st.session_state[f"delete_columns_{item['id']}"] = list(item["columns"])
    for item in st.session_state.pivot_tabs:
        st.session_state[f"pivot_category_{item['id']}"] = item["category"]
        st.session_state[f"pivot_value_{item['id']}"] = item["value"]
        st.session_state[f"pivot_split_{item['id']}"] = item.get("split")
        st.session_state[f"pivot_split_values_{item['id']}"] = list(
            item.get("split_selected", [])
        )


def is_missing(value):
    return bool(pd.isna(value))


def label_value(value):
    """Readable labels for Excel values, including missing cells."""
    return "(Blank / missing)" if is_missing(value) else str(value)


def possible_values(dataframe, column):
    """Show each distinct value from the entire original worksheet."""
    return dataframe[column].drop_duplicates().tolist()


def apply_filters(dataframe, filter_items):
    """OR within a filter's selected values; AND across different filters."""
    result = dataframe
    for item in filter_items:
        column = item["column"]
        selected_ids = item["selected"]
        if column is None or not selected_ids:
            continue

        options = possible_values(dataframe, column)
        chosen = [options[i] for i in selected_ids if 0 <= i < len(options)]
        ordinary_values = [value for value in chosen if not is_missing(value)]
        include_missing = any(is_missing(value) for value in chosen)

        mask = result[column].isin(ordinary_values)
        if include_missing:
            mask = mask | result[column].isna()
        result = result.loc[mask]
    return result


def columns_to_hide(delete_tabs):
    """A column selected in any Delete Column tab is hidden."""
    return {column for item in delete_tabs for column in item["columns"]}


def numeric_values(series):
    """Coerce ordinary and accounting-formatted numbers for summation.

    Values such as '$1,250', '1,250.00' and '(35.50)' are supported.
    Missing or nonnumeric inputs are ignored in the sum.
    """
    if pd.api.types.is_bool_dtype(series) or pd.api.types.is_datetime64_any_dtype(series):
        return pd.Series(float("nan"), index=series.index)
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce")

    cleaned = series.astype("string").str.strip()
    cleaned = cleaned.str.replace(r"^\((.*)\)$", r"-\1", regex=True)
    cleaned = cleaned.str.replace(r"[$£€¥,\s]", "", regex=True)
    return pd.to_numeric(cleaned, errors="coerce")


def numeric_columns(dataframe):
    """Offer numeric or numeric-text Excel fields as pivot measure options."""
    return [
        column for column in dataframe.columns
        if numeric_values(dataframe[column]).notna().any()
    ]


def make_pivot_table(filtered_data, category, value):
    """Sum a measure for each distinct category in the filtered worksheet."""
    source = pd.DataFrame(
        {category: filtered_data[category], value: numeric_values(filtered_data[value])}
    )
    return (
        source.groupby(category, dropna=False, sort=False)[value]
        .sum(min_count=1)
        .reset_index()
    )


def make_split_pivot_table(filtered_data, category, value, split, selected_values):
    """Group by category; make one sum column per selected split value.

    Group rows come from all filtered records. Missing combinations are 0,
    including categories whose selected split values never occur.
    The Total column sums ONLY the selected split categories.
    """
    source = pd.DataFrame(
        {
            "_group": filtered_data[category],
            "_split": filtered_data[split],
            "_measure": numeric_values(filtered_data[value]),
        },
        index=filtered_data.index,
    )
    group_index = source.groupby("_group", dropna=False, sort=False).size().index
    output = pd.DataFrame({category: group_index.tolist()})
    sum_columns = []
    used_labels = {str(category).casefold(), "total"}

    for split_value in selected_values:
        base_label = label_value(split_value)
        title = base_label
        suffix = 2
        while title.casefold() in used_labels:
            title = f"{base_label} ({suffix})"
            suffix += 1
        used_labels.add(title.casefold())

        if is_missing(split_value):
            mask = source["_split"].isna()
        else:
            mask = source["_split"].eq(split_value).fillna(False)

        totals = (
            source.loc[mask]
            .groupby("_group", dropna=False, sort=False)["_measure"]
            .sum(min_count=1)
        )
        output[title] = totals.reindex(group_index, fill_value=0).fillna(0).to_numpy()
        sum_columns.append(title)

    output["Total"] = output[sum_columns].sum(axis=1)
    return output


def unique_pivot_title(requested, existing):
    """Preserve full display tab titles, including ones longer than Excel allows."""
    existing_folded = {title.casefold() for title in existing}
    title = requested
    number = 2
    while title.casefold() in existing_folded:
        title = f"{requested} ({number})"
        number += 1
    return title


def excel_sheet_name(requested, existing):
    """Create a unique legal worksheet name (Excel max 31 characters)."""
    import re

    safe_name = re.sub(r"[\[\]:*?/\\]", "_", str(requested)).strip(" '")
    safe_name = safe_name or "Pivot"
    names_in_use = {name.casefold() for name in existing}
    name = safe_name[:31]
    suffix_id = 2
    while name.casefold() in names_in_use:
        suffix = f" ({suffix_id})"
        name = safe_name[: 31 - len(suffix)] + suffix
        suffix_id += 1
    return name


def pivot_results(filtered_data, pivot_tabs, original_data=None):
    """Build pivot result tabs from ALL filtered rows (not the displayed subset)."""
    result = []
    existing_names = ["Processed Data"]
    original_data = filtered_data if original_data is None else original_data

    for item in pivot_tabs:
        if item["category"] is None or item["value"] is None:
            continue
        category, value = item["category"], item["value"]
        split = item.get("split")
        if split is None:
            name = unique_pivot_title(f"Pivot {category}_{value}", existing_names)
            pivot_df = make_pivot_table(filtered_data, category, value)
        else:
            options = possible_values(original_data, split)
            selected_values = [
                options[i] for i in item.get("split_selected", [])
                if 0 <= i < len(options)
            ]
            if not selected_values:
                continue  # A split pivot is completed after selecting its categories.
            name = unique_pivot_title(
                f"Pivot {category}_{value}_Separated_by_{split}", existing_names
            )
            pivot_df = make_split_pivot_table(
                filtered_data, category, value, split, selected_values
            )
        existing_names.append(name)
        result.append((name, pivot_df))
    return result


def build_excel_file(processed_df, pivots=()):
    """Export all filtered rows, remaining columns and created pivot sheets."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        used_sheet_names = []
        for display_title, frame in [("Processed Data", processed_df), *pivots]:
            # Excel limits worksheet names to 31 characters. Streamlit tab titles
            # retain the user's full requested name; Excel uses a readable alias.
            sheet_title = display_title
            if len(sheet_title) > 31 and "_Separated_by_" in sheet_title:
                sheet_title = sheet_title.replace("_Separated_by_", "_by_", 1)
            sheet_name = excel_sheet_name(sheet_title, used_sheet_names)
            used_sheet_names.append(sheet_name)
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.sheets[sheet_name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

            for cell in sheet[1]:
                cell.fill = PatternFill(fill_type="solid", fgColor="183153")
                cell.font = Font(color="FFFFFF", bold=True)

            # Treat cells beginning with '=' as literal data, not formulas.
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type == "f" and isinstance(cell.value, str):
                        cell.data_type = "s"

            for column_index, column_name in enumerate(frame.columns, start=1):
                sample = frame[column_name].head(50)
                widths = [len(str(column_name))]
                widths.extend(len(str(value)) for value in sample if pd.notna(value))
                sheet.column_dimensions[get_column_letter(column_index)].width = min(
                    max(max(widths) + 2, 12), 42
                )
    return output.getvalue()


st.title("Accounting Excel Preview")
st.write("Upload an Excel workbook to filter, hide columns, build pivot summaries, and export the results.")

uploaded_file = st.file_uploader(
    "Choose an Excel file",
    type=["xlsx", "xls"],
    help="Supported file formats: .xlsx and .xls",
)

if uploaded_file is None:
    st.info("Upload an Excel file to get started.")
else:
    try:
        workbook = pd.ExcelFile(uploaded_file)
        sheet_name = st.selectbox("Choose a worksheet", workbook.sheet_names)
        df = pd.read_excel(workbook, sheet_name=sheet_name, header=0)
    except Exception as error:
        st.error(f"Could not read this Excel file: {error}")
    else:
        # Switching file or sheet starts a fresh set of filters and deleted columns.
        workbook_hash = hashlib.sha256(uploaded_file.getvalue()).hexdigest()
        source_id = (workbook_hash, sheet_name)
        if st.session_state.get("filter_source_id") != source_id:
            reset_workspace(source_id)
        else:
            prepare_workspace()

        # Apply row filters FIRST, then remove columns from the processed table.
        # Filters remain effective even when their source columns are hidden.
        filtered_df = apply_filters(df, st.session_state.filters)
        hidden = columns_to_hide(st.session_state.delete_tabs)
        processed_df = filtered_df.drop(columns=list(hidden), errors="ignore")
        pivots = pivot_results(filtered_df, st.session_state.pivot_tabs, df)

        st.success(f"Successfully loaded: {uploaded_file.name}")
        row_metric, source_cols_metric, shown_cols_metric = st.columns(3)
        row_metric.metric("Data rows", f"{len(df):,}")
        source_cols_metric.metric("Original columns", f"{len(df.columns):,}")
        shown_cols_metric.metric("Displayed columns", f"{len(processed_df.columns):,}")

        st.subheader("Column headers")
        if len(processed_df.columns):
            st.write(" | ".join(str(header) for header in processed_df.columns))
        else:
            st.caption("All columns are currently hidden from the preview.")

        function_label, filter_button, delete_button, export_button, back_button, spacer = st.columns(
            [1.6, 1.2, 2.1, 1.2, 1.2, 2.7]
        )
        with function_label:
            st.subheader("Function")
        with filter_button:
            st.button("Filter", on_click=add_filter, use_container_width=True)
        with delete_button:
            st.button(
                "Delete Column",
                on_click=add_delete_column,
                use_container_width=True,
                help="Create a tab to hide columns from the table preview.",
            )
        with export_button:
            # A download button exports the current processed table immediately.
            # It includes all matching rows, not only those visible in the scroller.
            if len(processed_df.columns) > 0:
                try:
                    export_bytes = build_excel_file(processed_df, pivots)
                except Exception as error:
                    st.error(f"Could not prepare the Excel export: {error}")
                else:
                    st.download_button(
                        "Export",
                        data=export_bytes,
                        file_name=f"{Path(uploaded_file.name).stem}_processed.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        help="Download the processed data and all created pivots as Excel sheets.",
                    )
            else:
                st.button(
                    "Export",
                    disabled=True,
                    use_container_width=True,
                    help="Keep at least one column to export an Excel file.",
                )
        with back_button:
            st.button(
                "↶ Back",
                on_click=go_back,
                disabled=not st.session_state.filter_history,
                help="Undo the latest filter, deletion, or pivot change.",
                use_container_width=True,
            )

        if st.session_state.function_tabs:
            # One tab bar for both functions, in the order the user added them.
            tab_titles = []
            filter_number = 0
            delete_number = 0
            for entry in st.session_state.function_tabs:
                if entry["kind"] == "filter":
                    filter_number += 1
                    tab_titles.append(f"Filter {filter_number}")
                else:
                    delete_number += 1
                    tab_titles.append(f"Delete Column {delete_number}")

            tab_containers = st.tabs(tab_titles)
            for tab, entry in zip(tab_containers, st.session_state.function_tabs):
                with tab:
                    if entry["kind"] == "filter":
                        item = next(
                            f for f in st.session_state.filters if f["id"] == entry["id"]
                        )
                        filter_id = item["id"]
                        st.selectbox(
                            "Column header",
                            options=[None] + list(df.columns),
                            format_func=lambda column: (
                                "Choose a column..." if column is None else str(column)
                            ),
                            key=f"filter_column_{filter_id}",
                            on_change=change_column,
                            args=(filter_id,),
                        )
                        if item["column"] is not None:
                            choices = possible_values(df, item["column"])
                            st.multiselect(
                                "Select one or more values to include",
                                options=list(range(len(choices))),
                                format_func=lambda i, values=choices: label_value(values[i]),
                                key=f"filter_values_{filter_id}",
                                on_change=change_values,
                                args=(filter_id,),
                                help="Leave blank to show all values for this column.",
                            )
                    else:
                        delete_id = entry["id"]
                        st.multiselect(
                            "Select columns to delete from the displayed table",
                            options=list(df.columns),
                            format_func=str,
                            key=f"delete_columns_{delete_id}",
                            on_change=change_deleted_columns,
                            args=(delete_id,),
                            help=(
                                "Selected columns will be removed from the preview and exported file. "
                                "Your uploaded Excel file and all filters remain unchanged."
                            ),
                        )
                        st.caption(
                            "Selected columns disappear from the preview and export. "
                            "Filters on those columns continue to work."
                        )

        # The pivot action has its own row, below Function and above results.
        create_label, pivot_button, pivot_spacer = st.columns([2.8, 1.4, 6.8])
        with create_label:
            st.subheader("Create New Tab")
        with pivot_button:
            st.button(
                "Pivot",
                on_click=add_pivot,
                use_container_width=True,
                help="Create a new pivot table that sums one numeric column by category.",
            )

        if st.session_state.pivot_tabs:
            available_numeric_columns = numeric_columns(df)
            st.caption(
                "Choose the group and numeric sum columns. Optionally use Split by "
                "to create a separate column for each selected category, plus a Total. "
                "Pivots use all filtered rows, even if a source column is hidden."
            )
            pivot_setup_tabs = st.tabs(
                [f"Pivot setup {i + 1}" for i in range(len(st.session_state.pivot_tabs))]
            )
            for setup_tab, config in zip(pivot_setup_tabs, st.session_state.pivot_tabs):
                pivot_id = config["id"]
                with setup_tab:
                    category_options = [None] + list(df.columns)
                    st.selectbox(
                        "Categorical column (group by)",
                        options=category_options,
                        format_func=lambda x: "Choose a categorical column..." if x is None else str(x),
                        key=f"pivot_category_{pivot_id}",
                        on_change=change_pivot_category,
                        args=(pivot_id,),
                    )
                    candidates = [
                        column for column in available_numeric_columns
                        if column != config["category"]
                    ]
                    st.selectbox(
                        "Continuous / numeric column (sum)",
                        options=[None] + candidates,
                        format_func=lambda x: "Choose a numeric column..." if x is None else str(x),
                        key=f"pivot_value_{pivot_id}",
                        on_change=change_pivot_value,
                        args=(pivot_id,),
                        disabled=config["category"] is None,
                    )
                    split_candidates = [
                        column for column in df.columns
                        if column != config["category"] and column != config["value"]
                    ]
                    st.selectbox(
                        "Split by (optional)",
                        options=[None] + split_candidates,
                        format_func=lambda x: "No split" if x is None else str(x),
                        key=f"pivot_split_{pivot_id}",
                        on_change=change_pivot_split,
                        args=(pivot_id,),
                        disabled=config["category"] is None or config["value"] is None,
                        help="Turn selected category values into separate sum columns.",
                    )
                    if config.get("split") is not None:
                        split_choices = possible_values(df, config["split"])
                        st.multiselect(
                            "Choose categories to include",
                            options=list(range(len(split_choices))),
                            format_func=lambda i, values=split_choices: label_value(values[i]),
                            key=f"pivot_split_values_{pivot_id}",
                            on_change=change_pivot_split_values,
                            args=(pivot_id,),
                            help="The Total column sums only these selected categories.",
                        )
                        if not config.get("split_selected"):
                            st.info("Select at least one category to create the split pivot.")
                        else:
                            st.success(
                                "Created pivot: "
                                f"Pivot {config['category']}_{config['value']}_Separated_by_{config['split']}"
                            )
                    elif config["category"] is not None and config["value"] is not None:
                        st.success(f"Created pivot: Pivot {config['category']}_{config['value']}")
                    elif not available_numeric_columns:
                        st.warning("No numeric columns found in the uploaded worksheet.")

        # Result tabs match the sheets in the exported Excel workbook.
        st.subheader("Results")
        results_tabs = st.tabs(["Processed Data"] + [name for name, _ in pivots])
        with results_tabs[0]:
            st.caption(
                f"{len(filtered_df):,} matching rows out of {len(df):,} total; "
                f"{len(processed_df.columns):,} columns remaining. "
                "Scroll inside the table to see all matching rows (about 10 at a time). "
                "Export includes every matching row."
            )
            if filtered_df.empty:
                st.warning("No rows match the filters. The export will contain headers only.")
            if processed_df.shape[1] == 0:
                st.info("All columns are hidden. Deselect a column or click Back to restore it.")
            else:
                st.dataframe(
                    processed_df,
                    height=395,
                    use_container_width=True,
                    hide_index=True,
                )

        for result_tab, (tab_title, pivot_df) in zip(results_tabs[1:], pivots):
            with result_tab:
                st.caption(
                    f"{len(pivot_df):,} categories. Sums reflect the current row filters. "
                    "Scroll to see all categories; Export includes this sheet."
                )
                st.dataframe(
                    pivot_df,
                    height=395,
                    use_container_width=True,
                    hide_index=True,
                )
